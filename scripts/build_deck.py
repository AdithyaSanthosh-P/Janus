"""Build the 8-slide submission deck from the official template.

    python scripts/build_deck.py --team "Team Name" --college "College Name" \
        --members "A <a@x.com>; B <b@x.com>" --repo https://github.com/... \
        [--out docs/deck/Janus_Submission.pptx]

Keeps template slides 1, 2, 4, 5, 8, 7, 6, 9 (title, problem, architecture,
an interruption handled, results, extension, stack and rigour, what's next) in
that order and drops the rest, so the deck stays within the guide's 8 slides.
Dev-time only (needs python-pptx); the committed .pptx is the artifact.
"""

from __future__ import annotations

import argparse
import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "guidelines" / "mail" / "CollegeName_TeamName_Submission.pptx"
KEEP = [0, 1, 3, 4, 7, 6, 5, 8]  # template slide indexes, in deck order

INK = RGBColor(0x1B, 0x1F, 0x24)
BLUE, AMBER, GREEN, VIOLET, GREY = (RGBColor(0xDB, 0xEA, 0xFE), RGBColor(0xFE, 0xF3, 0xC7), RGBColor(0xDC, 0xFC, 0xE7),
                                    RGBColor(0xED, 0xE9, 0xFE), RGBColor(0xF1, 0xF3, 0xF5))


def set_text(shape, lines, size=18, bold_first=False):
    tf = shape.text_frame
    tf.word_wrap = True
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bool(bold_first and i == 0)
    return tf


def body_box(slide, x, y, w, h):
    ph = slide.placeholders[1]
    ph.left, ph.top, ph.width, ph.height = Inches(x), Inches(y), Inches(w), Inches(h)
    return ph


def title(slide, text):
    t = slide.placeholders[0]
    set_text(t, [text], size=30, bold_first=True)
    t.top, t.height = Inches(0.35), Inches(1.0)


def box(slide, x, y, w, h, text, fill=GREY, size=12):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
    tf = s.text_frame
    tf.word_wrap = True
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size if i == 0 else size - 2)
        r.font.bold = i == 0
        r.font.color.rgb = INK
    return s


def arrow(slide, x1, y1, x2, y2):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    c.line.width = Pt(1.75)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.append(tail)


def table(slide, x, y, w, rows, col_w, size=13):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.42 * len(rows)))
    t = shape.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(size)
                    r.font.bold = i == 0 or (i == 1 and True)
    return t


def build(args) -> Presentation:
    prs = Presentation(str(TEMPLATE))
    slides = list(prs.slides)

    # 1. title -------------------------------------------------------------
    s = slides[0]
    info = next(sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.startswith("Theme ID"))
    set_text(info, [
        "Theme ID - Theme 05: Interruptible Real-Time Agents",
        f"Team Name - {args.team}",
        f"College Name - {args.college}",
        *[f"Member - {m.strip()}" for m in args.members.split(";") if m.strip()],
        f"Submission Github link - {args.repo}",
    ], size=14)

    # 2. problem -------------------------------------------------------------
    s = slides[1]
    title(s, "Problem: an agent that acts is unsafe to interrupt")
    set_text(body_box(s, 0.9, 1.6, 10.6, 4.8), [
        "People change their mind mid-sentence: “book Thursday … no wait, Friday”.",
        "An LLM in a loop books twice, books the old value, or announces success it did not achieve.",
        "Full-Duplex-Bench v3 tests exactly this over real audio: disfluent speech, self-corrections, multi-step tasks, and it scores strict Pass@1. The best published system passes 60%.",
        "Our aim: interruption-safe by construction, not by prompting.",
    ], size=22)

    # 3. architecture ------------------------------------------------------------
    s = slides[3 - 1 + 1]  # template slide 4 (index 3)
    s = slides[3]
    title(s, "Solution: models propose, one kernel decides")
    box(s, 0.5, 1.7, 1.7, 0.9, "Audio in\nSilero VAD + Whisper", BLUE)
    box(s, 0.5, 4.7, 1.7, 0.9, "Audio out\nKokoro TTS", BLUE)
    box(s, 2.9, 1.5, 2.6, 2.2, "Janus kernel\none synchronous step,\nsole writer of state\nread-set invalidation\nsame-step cancel", AMBER, 13)
    box(s, 2.9, 4.2, 2.6, 1.0, "CommitGate + EmissionGate\nsettle · intent · no duplicate", AMBER)
    box(s, 6.1, 1.5, 2.0, 1.1, "Async workers\nInterpret · Plan\nCompose · Vision", VIOLET)
    box(s, 6.1, 3.0, 2.0, 0.8, "Gemini\n(hosted)", VIOLET)
    box(s, 6.1, 4.3, 2.0, 0.9, "Tools\nFDB mock APIs /\ndevice care", GREEN)
    arrow(s, 2.2, 2.15, 2.9, 2.4)
    arrow(s, 4.2, 3.7, 4.2, 4.2)
    arrow(s, 2.9, 4.7, 2.2, 5.0)
    arrow(s, 5.5, 2.0, 6.1, 2.0)
    arrow(s, 6.1, 2.35, 5.5, 2.6)
    arrow(s, 7.1, 2.6, 7.1, 3.0)
    arrow(s, 5.5, 4.7, 6.1, 4.7)
    set_text(body_box(s, 8.4, 1.6, 3.3, 4.6), [
        "Workers only return proposals; the kernel never awaits.",
        "Every proposal records the facts it read. A changed fact cancels dependent work in the same step; a reverted value costs nothing.",
        "A write passes one gate and is ledgered before it is emitted: no double booking.",
    ], size=16)

    # 4. interruption handled --------------------------------------------------------
    s = slides[4]
    title(s, "An interruption, handled (kernel X-ray)")
    s.shapes.add_picture(str(ROOT / "docs" / "xray" / "correct_a_running_lookup.png"), Inches(0.5), Inches(1.5), width=Inches(7.3))
    set_text(body_box(s, 8.0, 1.6, 3.8, 4.8), [
        "User asks about an orange light; 2.4 s later: “sorry, actually red”.",
        "The orange lookup is struck through: invalidated and cancelled in the same step the correction arrives.",
        "Re-run with red; one final answer grounded in the red result.",
    ], size=16)

    # 5. results --------------------------------------------------------------------
    s = slides[7]
    title(s, "Results on Full-Duplex-Bench v3")
    set_text(body_box(s, 0.9, 5.05, 10.6, 1.6), [
        "Text replay feeds ground-truth transcripts: it measures the reasoning and safety core, not speech or latency. Voice row: all 100 recordings on the final code (3 Oct); three full voice runs that day scored 73, 65 and 74.",
        "Housing is weakest (35% text), as for every published system; 17 of the 26 voice misses are labels that disagree with the published tools or the recording (README: Benchmark labels). The organizers' re-run of ./reproduce.sh is what scores.",
    ], size=14)
    table(s, 0.9, 1.6, 10.6, [
        ["System (FDB-v3, gpt-4o judge)", "Pass@1", "Tool sel.", "Arg acc.", "Resp. qual."],
        ["Janus, text replay, 100 scenarios", "77.0%", "0.959", "0.807", "0.840"],
        ["Janus, voice over LiveKit, 100 (3 Oct, final)", "74.0%", "0.959", "0.780", "0.720"],
        ["GPT-Realtime (paper)", "60.0%", "0.876", "0.680", "0.792"],
        ["Gemini Live 3.1 (paper)", "54.0%", "0.817", "0.588", "0.718"],
        ["Cascaded Whisper→GPT-4o→TTS (paper)", "45.0%", "0.803", "0.562", "0.600"],
    ], [4.6, 1.5, 1.5, 1.5, 1.5])

    # 6. extension --------------------------------------------------------------------
    s = slides[6]
    title(s, "Extension: Samsung-style home care")
    s.shapes.add_picture(str(ROOT / "docs" / "xray" / "withdraw_one_of_two_writes.png"), Inches(0.5), Inches(1.5), width=Inches(7.3))
    set_text(body_box(s, 8.0, 1.5, 3.9, 5.0), [
        "Router read from the camera; a Samsung washer (codes 4C, 5C, UB, dC) through a SmartThings-shaped status.",
        "“Stop the wash … wait, just pause it”: the pending stop is never sent; pause and the booking go through once.",
        "“Resume” while 4C shows: the washer refuses, and the agent says so.",
        "Effect ledger per session: committed, withdrawn, refused, duplicates (0).",
        "Real: speech, camera, Gemini, the kernel. Simulated: the device and Samsung Care services.",
    ], size=14)

    # 7. rigour / stack / limits -----------------------------------------------------------
    s = slides[5]
    title(s, "Stack, rigour and honest limits")
    set_text(body_box(s, 0.9, 1.6, 10.6, 4.9), [
        "Stack: LiveKit Agents · Silero VAD · faster-whisper large-v3-turbo · Kokoro-82M · Gemini 3.6 Flash (hosted) · Python 3.11, Docker.",
        "645 deterministic tests (stepped clock, scripted model), TraceChecker on every run, adversarial explorer over 19 race scenarios.",
        "One-command reproduction: pinned FDB commit, sha256-checked data, weights baked into the image.",
        "Limits: Janus follows the published tool contract, not labels that disagree with it (e.g. an undeclared pets_allowed); housing 35%; the hosted model is not bit-reproducible; the extension's device backend is simulated.",
    ], size=20)

    # 8. what's next -----------------------------------------------------------------------------
    s = slides[8]
    title(s, "What's next")
    set_text(body_box(s, 0.9, 1.6, 10.6, 4.9), [
        "A live SmartThings adapter behind the same tool manifest (the demo's device backend is simulated).",
        "Measure the full 100-recording voice path on organizer-class hardware and close the housing gap.",
        "Samsung worklet: one safe action layer for any device manifest, so a phone, TV or appliance gets interruption-safe voice control from its tool list alone.",
        "On-device speech and a small local decision model behind the same kernel.",
    ], size=21)

    # keep only the chosen slides, in order
    sld_ids = prs.slides._sldIdLst
    entries = list(sld_ids)
    for i, e in enumerate(entries):
        if i not in KEEP:
            prs.part.drop_rel(e.rId)
            sld_ids.remove(e)
    kept = [e for i, e in enumerate(entries) if i in KEEP]
    for e in kept:
        sld_ids.remove(e)
    for i in KEEP:
        sld_ids.append(entries[i])
    return prs


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--team", default="[TEAM NAME]")
    ap.add_argument("--college", default="[COLLEGE NAME]")
    ap.add_argument("--members", default="[Member 1 name & email]; [Member 2 name & email]")
    ap.add_argument("--repo", default="[public GitHub URL]")
    ap.add_argument("--out", default=str(ROOT / "docs" / "deck" / "Janus_Submission.pptx"))
    args = ap.parse_args()
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    build(args).save(out)
    print("wrote", out)
